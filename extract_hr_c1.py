#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Extractor del tablero HORA A HORA - CELDA C1.
(Infusion / Xcela / Smart Port / Boxing)

Distinto al hora-a-hora de AngioDynamics: aqui NO hay una tabla resumida con
metas ya calculadas -- cada hoja semanal (WWnn) trae directo la rejilla de
horas con columnas "Real" y "Scrap" por cada turno, y dentro de cada dia hay
4 categorias (Infusion, Xcela, Smart Port, Boxing), cada una con 2 filas:
una fila "etiqueta" (solo el nombre) y la fila de abajo con los datos.

Estructura de cada hoja WWnn:
  Fila 1: encabezado de cada turno, ej. "TURNO A   ·   6:00 am - 3:36 pm"
  Fila 2: etiqueta de cada hora (cada hora ocupa 2 columnas), y "TOTAL TURNO"
          al final de cada bloque de turno.
  Fila 3: "Real" / "Scrap" bajo cada hora.
  Columna A: marca el dia ("  LUNES -- ") y debajo las categorias.

No hay meta/UPH definida todavia (pendiente, se agregara despues) -- por eso
esto NO calcula status de cumplimiento, solo expone los numeros tal cual.

USO:
  python extract_hr_c1.py            (una vez)
  python extract_hr_c1.py --watch    (cada 15 min)
"""
import json
import re
import sys
import time
import unicodedata
import datetime
from pathlib import Path

from openpyxl import load_workbook

# ----------------------------------------------------------------------------
# CONFIGURACION -- AJUSTA ESTA RUTA A TU ARCHIVO REAL
# ----------------------------------------------------------------------------
EXCEL_PATH = Path(
    r"C:\Users\aaron.lara\OneDrive - Biomerics\BALA-CENTRAL - 1. Producción\1-NEW FILES\C1\Hr_Hr_C1 - (1).xlsx"
)
OUTPUT_JSON = Path(__file__).parent / "data_hr_c1.json"
REFRESH_SECONDS = 900  # 15 min

DIAS = ["LUNES", "MARTES", "MIERCOLES", "JUEVES", "VIERNES", "SABADO", "DOMINGO"]
HEADER_ROW_TURNOS = 1
HEADER_ROW_HORAS = 2
HEADER_ROW_REALSCRAP = 3
DATA_START_ROW = 4


def _quitar_acentos(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def _num(v):
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        if s == "" or s.startswith("#"):
            return None
        try:
            return float(s)
        except ValueError:
            return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def _clean(s):
    return re.sub(r"\s+", " ", str(s).strip())


def find_turnos(ws):
    """Ubica cada bloque de turno (A/B/C) leyendo la fila 1, y dentro de cada
    uno arma la lista de horas (par de columnas Real/Scrap) hasta topar con
    la columna 'TOTAL TURNO'."""
    starts = []
    for c in range(1, ws.max_column + 1):
        raw = ws.cell(HEADER_ROW_TURNOS, c).value
        if isinstance(raw, str) and "TURNO" in raw.upper():
            m = re.search(r"TURNO\s+([A-Z])", raw.upper())
            letra = m.group(1) if m else str(len(starts) + 1)
            starts.append({"letra": letra, "col_inicio": c})

    turnos = []
    for i, blk in enumerate(starts):
        col_fin = (starts[i + 1]["col_inicio"] - 1) if i + 1 < len(starts) else ws.max_column
        hour_labels = []
        hour_cols = []  # (col_real, col_scrap)
        total_real_col = None
        total_scrap_col = None
        c = blk["col_inicio"]
        while c <= col_fin:
            lbl = ws.cell(HEADER_ROW_HORAS, c).value
            if isinstance(lbl, str) and lbl.strip():
                if lbl.strip().upper().startswith("TOTAL"):
                    total_real_col = c
                    total_scrap_col = c + 1
                    break
                hour_labels.append(_clean(lbl))
                hour_cols.append((c, c + 1))
            c += 1
        turnos.append({
            "letra": blk["letra"],
            "hour_labels": hour_labels,
            "hour_cols": hour_cols,
            "total_real_col": total_real_col,
            "total_scrap_col": total_scrap_col,
        })
    return turnos


def find_day_rows(ws):
    rows = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str):
            up = _quitar_acentos(v.strip().upper())
            for d in DIAS:
                if d in up and d not in rows:
                    rows[d] = r
    return rows


def read_dia(ws, fila_inicio, fila_fin, turnos):
    """Lee las categorias (Infusion/Xcela/Smart Port/Boxing) de un dia.
    Una fila es 'de datos' si tiene algo en la columna Total Real del
    PRIMER turno (las filas de puro rotulo, arriba de cada categoria,
    vienen vacias ahi)."""
    ancla = turnos[0]["total_real_col"] if turnos else None
    if ancla is None:
        return []

    procesos = []
    for r in range(fila_inicio, fila_fin + 1):
        if _num(ws.cell(r, ancla).value) is None:
            continue
        nombre = ws.cell(r, 1).value
        if not nombre or not str(nombre).strip():
            continue

        por_turno = {}
        for t in turnos:
            real = [_num(ws.cell(r, cr).value) for cr, cs in t["hour_cols"]]
            scrap = [_num(ws.cell(r, cs).value) for cr, cs in t["hour_cols"]]
            total_real = _num(ws.cell(r, t["total_real_col"]).value) if t["total_real_col"] else None
            total_scrap = _num(ws.cell(r, t["total_scrap_col"]).value) if t["total_scrap_col"] else None
            por_turno[t["letra"]] = {
                "hour_labels": t["hour_labels"],
                "real": real,
                "scrap": scrap,
                "total_real": total_real if total_real is not None else sum(v or 0 for v in real),
                "total_scrap": total_scrap if total_scrap is not None else sum(v or 0 for v in scrap),
            }
        procesos.append({"name": _clean(nombre), "turnos": por_turno})
    return procesos


def read_sheet(ws):
    turnos = find_turnos(ws)
    if not turnos:
        return None
    day_rows = find_day_rows(ws)
    if not day_rows:
        return None
    ordenados = sorted(day_rows.items(), key=lambda kv: kv[1])
    dias = {}
    for i, (dia, fila) in enumerate(ordenados):
        fin = (ordenados[i + 1][1] - 1) if i + 1 < len(ordenados) else ws.max_row
        procesos = read_dia(ws, fila + 1, fin, turnos)
        if procesos:
            dias[dia] = procesos
    return dias if dias else None


def build_payload():
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(f"No se encontro el Excel en: {EXCEL_PATH}")

    wb = load_workbook(EXCEL_PATH, data_only=True)
    weeks = {}
    week_order = []
    for name in wb.sheetnames:
        if not re.match(r"^W{1,2}\d+$", name.strip(), flags=re.I):
            continue
        dias = read_sheet(wb[name])
        weeks[name] = {"label": name, "dias": dias or {}}
        week_order.append(name)

    payload = {
        "plant": "Celda C1",
        "source_file": EXCEL_PATH.name,
        "generated_at": datetime.datetime.now().isoformat(timespec="seconds"),
        "weeks": weeks,
        "week_order": week_order,
        "dias_orden": DIAS,
    }
    return payload


def write_json():
    p = build_payload()
    OUTPUT_JSON.write_text(json.dumps(p, ensure_ascii=False, indent=1), encoding="utf-8")
    con_datos = sum(1 for w in p["weeks"].values() if w["dias"])
    print(f"[{datetime.datetime.now():%H:%M:%S}] data_hr_c1.json: "
          f"{len(p['week_order'])} semanas ({con_datos} con datos)")


def main():
    watch = "--watch" in sys.argv
    while True:
        try:
            write_json()
        except Exception as e:
            print(f"[ERROR] {e}", file=sys.stderr)
        if not watch:
            break
        time.sleep(REFRESH_SECONDS)


if __name__ == "__main__":
    main()
